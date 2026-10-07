from pptx import Presentation
p = Presentation("IntelliScrap_Presentation.pptx")
print("slides:", len(p.slides))
for i, s in enumerate(p.slides, 1):
    texts = []
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            texts.append(sh.text_frame.text.strip().splitlines()[0])
    print(f"[{i:2}] {'; '.join(texts[:3])[:110]}")
